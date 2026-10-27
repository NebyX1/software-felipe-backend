from io import BytesIO
import zipfile
from flask import request,g,send_file,current_app
from ..extensions import db
from ..models import User,Unit,ActivityLog
from ..models.domain import Company,Professional,CatalogEntry,Intervention,History,Contract,BudgetLine,Document
from ..services.access import authenticated,roles,work_for,work_query,documents_query,history
from ..services.errors import require
from ..services.works import text
from ..services.reports import pdf_report,csv_report
from ..infrastructure.storage import get_document_storage
from .interventions import body,filtered
from . import api_bp


@api_bp.get('/reference')
@authenticated
def reference():
    companies=Company.query.filter_by(id=g.user.company_id).all() if g.user.role=='empresa' else Company.query.order_by(Company.name).all()
    return {'companies':[c.to_dict() for c in companies],'units':[u.to_dict() for u in Unit.query],'catalog':[c.to_dict() for c in CatalogEntry.query.filter_by(active=True)],'professionals':[] if g.user.role=='empresa' else [{'id':p.id,'name':p.name} for p in Professional.query], 'users':[] if g.user.role=='empresa' else [{'id':u.id,'name':u.name,'role':u.role,'unit_id':u.unit_id} for u in User.query.filter_by(is_active=True).filter(User.role!='empresa')]}


@api_bp.route('/admin/<kind>',methods=['GET','POST'])
@authenticated
def administration(kind):
    roles('super_admin','admin')
    models={'companies':Company,'professionals':Professional,'catalog':CatalogEntry,'users':User,'units':Unit}
    require(kind in models,'Administración no encontrada.',404)
    model=models[kind]
    if request.method=='GET':
        return {'items':[r.to_dict() for r in model.query.order_by(model.id)]}
    data=body()
    require(isinstance(data.get('contact',{}),dict) and isinstance(data.get('data',{}),dict),'Contactos y contexto deben ser objetos JSON.')
    if kind=='companies':
        row=Company(name=text(data.get('name',''),'Empresa',200),rut=text(data.get('rut',''),'RUT',40),contact=data.get('contact',{}),status='pendiente_validacion')
    elif kind=='professionals':
        row=Professional(name=text(data.get('name',''),'Nombre',200),registration=text(data.get('registration',''),'Registro',100),contact=data.get('contact',{}),source=data.get('source','Alta manual'))
    elif kind=='catalog':
        require(data.get('kind') in ('objetivo','tipo','financiacion','unidad','rubro','precio','capa_geografica','parametro'),'Tipo de catálogo inválido.')
        if data['kind']=='precio':
            require(data.get('data',{}).get('source') and data['data'].get('date'),'Precio requiere fuente y fecha.')
        row=CatalogEntry(kind=data['kind'],code=text(data.get('code',''),'Código',120),label=text(data.get('label',''),'Nombre',250),data=data.get('data',{}))
    elif kind=='units':
        row=Unit(name=text(data.get('name',''),'Unidad',100),color=data.get('color','#136b85'),emoji=data.get('emoji',''))
    else:
        require(data.get('role') in ('tecnico','director_obra','inspector','administrativo','consulta','empresa','direccion_general'),'Rol inválido.')
        row=User(name=text(data.get('name',''),'Nombre',100),email=text(data.get('email',''),'Email',150),role=data['role'],unit_id=data.get('unit_id'),company_id=data.get('company_id'),demo=False,is_active=True)
        require(row.role!='empresa' or db.session.get(Company,row.company_id),'El representante necesita empresa válida.')
        row.set_password(text(data.get('password',''),'Contraseña ficticia',100))
        if data.get('demo') is True:
            require(current_app.config['DEMO_ENVIRONMENT'] and current_app.config['DEMO_AUTH_ENABLED'],'Cuenta demo fuera de entorno demo.')
            row.demo=True
    db.session.add(row)
    db.session.flush()
    if kind == 'users' and row.role == 'empresa':
        from datetime import date
        from ..models.domain import CompanyRepresentative
        db.session.add(CompanyRepresentative(company_id=row.company_id, user_id=row.id, status='pendiente', start_date=date.today()))
    db.session.add(ActivityLog(user_id=g.user.id,action='administracion:alta:'+kind))
    db.session.commit()
    return {'item':row.to_dict()},201


@api_bp.patch('/admin/<kind>/<int:entity_id>')
@authenticated
def edit_admin(kind,entity_id):
    roles('super_admin','admin')
    models={'companies':Company,'professionals':Professional,'catalog':CatalogEntry,'users':User,'units':Unit}
    require(kind in models,'Recurso no encontrado.',404)
    row=db.session.get(models[kind],entity_id)
    require(row,'Recurso no encontrado.',404)
    data=body()
    allowed={'companies':('name','contact','status'),'professionals':('name','contact','source'),'catalog':('label','data','active'),'users':('name','unit_id','company_id','is_active'),'units':('name','color','emoji')}[kind]
    for key in allowed:
        if key in data:
            if key=='status':
                require(data[key] in ('pendiente_validacion','activa','suspendida','baja'),'Estado de empresa inválido.')
                transitions={'pendiente_validacion':{'activa','baja'},'activa':{'suspendida','baja'},'suspendida':{'activa','baja'},'baja':{'pendiente_validacion'}}
                require(data[key]==row.status or data[key] in transitions[row.status],'Transición de empresa inválida.',409)
            if key in ('active','is_active'):require(isinstance(data[key],bool),'Activación debe ser booleana.')
            if key in ('contact','data'):require(isinstance(data[key],dict),'Contexto/contacto debe ser objeto JSON.')
            setattr(row,key,data[key])
    db.session.add(ActivityLog(user_id=g.user.id,action='administracion:editar:'+kind,details=str(entity_id)))
    db.session.commit()
    return {'item':row.to_dict()}


@api_bp.get('/admin/audit/events')
@authenticated
def audit():
    roles('super_admin','admin','direccion_general')
    page=max(1,request.args.get('page',1,type=int))
    rows=History.query.order_by(History.id.desc()).paginate(page=page,per_page=100,error_out=False)
    return {'events':[h.to_dict() for h in rows.items],'total':rows.total,'auth_events':[a.to_dict() for a in ActivityLog.query.order_by(ActivityLog.id.desc()).limit(100)]}


@api_bp.get('/reports/works.<format>')
@authenticated
def works_report(format):
    require(format in ('csv','xlsx','pdf'),'Formato de reporte inválido.')
    works=filtered(work_query()).all()
    headers=['Código','Intervención','Localidad','Etapa','Estado','Superficie m2','Presupuesto UYU']
    rows=[[w.code,w.name,w.locality,w.stage,w.status,str(w.surface),str(w.budget)] for w in works]
    return export_table(format,'obras',headers,rows)


def export_table(format,name,headers,rows):
    if format=='csv':
        content,mime=csv_report(headers,rows),'text/csv'
    elif format=='pdf':
        from datetime import datetime,timezone
        content,mime=pdf_report(name.capitalize(),['Generado '+datetime.now(timezone.utc).isoformat()],headers,rows),'application/pdf'
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        workbook=Workbook()
        sheet=workbook.active
        sheet.title=name[:31]
        sheet.append(headers)
        for row in rows:
            sheet.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='183A56')
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(55,max(15,max(len(str(c.value or '')) for c in col)+2))
        output=BytesIO()
        workbook.save(output)
        content,mime=output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    db.session.add(ActivityLog(user_id=g.user.id,action='exportar:'+name+':'+format))
    db.session.commit()
    return send_file(BytesIO(content),mimetype=mime,as_attachment=True,download_name=name+'.'+format)


@api_bp.get('/works/<int:work_id>/contracts/<int:contract_id>/export.<format>')
@authenticated
def rubrado_export(work_id,contract_id,format):
    work_for(work_id)
    require(format in ('csv','xlsx','pdf'),'Formato inválido.')
    require(Contract.query.filter_by(id=contract_id,work_id=work_id).first(),'Contrato no encontrado.',404)
    headers=['code','chapter','description','unit','quantity','unit_price','social_rate','kind']
    rows=[[getattr(l,k) for k in headers] for l in BudgetLine.query.filter_by(contract_id=contract_id)]
    return export_table(format,'rubrado',headers,rows)


@api_bp.get('/works/<int:work_id>/dossier.zip')
@authenticated
def dossier(work_id):
    work=work_for(work_id)
    output=BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        for document in documents_query(work_id).filter(Document.status!='eliminado'):
            import hashlib
            content=get_document_storage().read(document.storage_key)
            require(hashlib.sha256(content).hexdigest()==document.sha256,'Documento alterado; expediente no exportado.',409)
            archive.writestr(f'{document.id}-v{document.version}-{document.name.replace(chr(92),"_").replace("/","_")}',content)
        import json
        archive.writestr('ficha-demo.json',json.dumps(work.to_dict(),ensure_ascii=False,indent=2))
    history(work_id,'obra',work_id,'exportar_expediente',details={'demo':True})
    db.session.commit()
    return send_file(BytesIO(output.getvalue()),mimetype='application/zip',as_attachment=True,download_name=work.code+'-expediente-demo.zip')
